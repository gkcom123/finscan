"""Pure-function tests for excel/cumulative.py's cumulative-to-standalone-quarter fix.

No graph, no LLM — a real (temporary) workbook stands in for "the prior quarters already
on disk", exactly the thing a real target model would have.
"""
from __future__ import annotations

from openpyxl import Workbook

from finscan.excel.cumulative import resolve_cumulative_periods
from finscan.excel.discovery import RowPlan, SheetPlan, WorkbookPlan
from finscan.schemas import PeriodMeta

FIELD = "depreciation_amortisation"
ROW = 10


def _write_workbook(tmp_path, prior_values: dict[int, float], sheet_name: str = "Model"):
    """A workbook with `sheet_name`'s row ROW holding prior_values at the given columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.cell(1, 1).value = "Label"
    ws.cell(ROW, 1).value = "Depreciation & Amortisation"
    for col, val in prior_values.items():
        ws.cell(ROW, col).value = val
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path


def _sheet_plan(
    *, write_col: int, populated_period_cols: list[int],
    period_dates: dict[int, str] | None = None, units: str = "units",
    sheet: str = "Model", score: float = 1.0,
) -> SheetPlan:
    return SheetPlan(
        sheet=sheet, in_scope=True, write_col=write_col, units=units, score=score,
        period_headers={}, period_dates=period_dates or {}, sections={},
        rows=[RowPlan(row=ROW, label="Depreciation & Amortisation", field=FIELD,
                      reference_role="input")],
        populated_period_cols=populated_period_cols,
    )


def _plan(path, *sheets: SheetPlan) -> WorkbookPlan:
    return WorkbookPlan(path=str(path), fingerprint="test", sheets=list(sheets), colors_found=True)


def _meta(period_type: str = "quarter") -> PeriodMeta:
    return PeriodMeta(period_label="Q2 FY2027", period_end_date="2026-09-30",
                      period_type=period_type, units="units")


def test_q1_is_a_true_no_op():
    """months_covered == this model's own single-period length -> nothing to subtract,
    falls out of the general formula with no special-cased branch."""
    sheet = _sheet_plan(write_col=2, populated_period_cols=[])
    values, issues = resolve_cumulative_periods(
        _plan("unused.xlsx", sheet), "unused.xlsx", {FIELD: 89.20}, {FIELD: 3}, None, _meta(),
    )
    assert values[FIELD] == 89.20
    assert not any(i.field == FIELD for i in issues)


def test_q2_subtracts_the_single_prior_quarter(tmp_path):
    sheet = _sheet_plan(write_col=3, populated_period_cols=[2])
    path = _write_workbook(tmp_path, {2: 100.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 250.0}, {FIELD: 6}, None, _meta(),
    )
    assert values[FIELD] == 150.0
    assert any(i.code == "cumulative_period_corrected" and i.field == FIELD for i in issues)


def test_q3_sums_two_prior_quarters_not_just_the_reference_column(tmp_path):
    """The general N-quarter case: Q3's 9-month cumulative minus Q1+Q2, proving this
    isn't hardcoded to only ever subtract a single reference column."""
    sheet = _sheet_plan(write_col=4, populated_period_cols=[2, 3])
    path = _write_workbook(tmp_path, {2: 100.0, 3: 110.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 360.0}, {FIELD: 9}, None, _meta(),
    )
    assert values[FIELD] == 360.0 - (100.0 + 110.0)
    assert any(i.code == "cumulative_period_corrected" and i.field == FIELD for i in issues)


def test_insufficient_history_removes_the_field_rather_than_guess(tmp_path):
    """A Q2 filing pointed at a workbook with no Q1 column yet: there is nothing to
    subtract, so the raw cumulative figure must NOT be written."""
    sheet = _sheet_plan(write_col=2, populated_period_cols=[])
    path = _write_workbook(tmp_path, {})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 250.0}, {FIELD: 6}, None, _meta(),
    )
    assert FIELD not in values
    assert any(i.code == "cumulative_period_insufficient_history" and i.severity == "error"
              and i.field == FIELD for i in issues)


def test_ambiguous_prior_column_dates_are_refused_not_guessed(tmp_path):
    """Two prior columns (2, 3) exist, but their own dates are 7 months apart, while the
    model's broader column history (cols 1 and 4 included) establishes a clear 3-month
    cadence — so this specific pair is the anomaly, not the norm, and summing them would
    be a guess rather than an honest read of "last quarter's own standalone figure"."""
    sheet = _sheet_plan(
        write_col=4, populated_period_cols=[2, 3],
        period_dates={1: "2025-10-31", 2: "2026-01-31", 3: "2026-08-31", 4: "2026-11-30"},
    )
    path = _write_workbook(tmp_path, {2: 100.0, 3: 110.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 360.0}, {FIELD: 9}, None, _meta(),
    )
    assert FIELD not in values
    assert any(i.code == "cumulative_period_ambiguous_history" and i.severity == "error"
              and i.field == FIELD for i in issues)


def test_implausible_correction_is_flagged_but_still_written(tmp_path):
    """A correction larger in magnitude than the raw cumulative figure it came from is
    a red flag (a standalone quarter should normally be smaller), but is written anyway
    with a warning, unlike insufficient/ambiguous history which are hard refusals."""
    sheet = _sheet_plan(write_col=3, populated_period_cols=[2])
    path = _write_workbook(tmp_path, {2: 200.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 50.0}, {FIELD: 6}, None, _meta(),
    )
    assert values[FIELD] == 50.0 - 200.0
    assert any(i.code == "implausible_cumulative_correction" and i.severity == "warning"
              for i in issues)


def test_cadence_mismatch_is_not_forced_into_a_nonsensical_n(tmp_path):
    """months=5 against a 3-month model cadence is not a clean multiple of quarters —
    don't force a wrong N, leave the value as extracted and warn instead."""
    sheet = _sheet_plan(write_col=3, populated_period_cols=[2])
    path = _write_workbook(tmp_path, {2: 100.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 250.0}, {FIELD: 5}, None, _meta(),
    )
    assert values[FIELD] == 250.0
    assert any(i.code == "cumulative_period_cadence_mismatch" for i in issues)


def test_multi_sheet_authority_is_computed_once_and_noted(tmp_path):
    """The same field mapped as a writable row on two sheets: one is used to compute
    the correction, and the choice is recorded for the reviewer to see."""
    primary = _sheet_plan(write_col=3, populated_period_cols=[2], sheet="Model", score=2.0)
    other = _sheet_plan(write_col=3, populated_period_cols=[2], sheet="Old model", score=1.0)
    path = _write_workbook(tmp_path, {2: 100.0}, sheet_name="Model")
    # "Old model" sheet must also exist in the same physical workbook for a fully
    # realistic fixture, even though only the higher-scored sheet's history is read.
    wb = __import__("openpyxl").load_workbook(path)
    wb.create_sheet("Old model")
    wb.save(path)
    plan = _plan(path, primary, other)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 250.0}, {FIELD: 6}, None, _meta(),
    )
    assert values[FIELD] == 150.0
    note = next(i for i in issues if i.code == "cumulative_period_multi_sheet_authority")
    assert "Model" in note.message and "Old model" in note.message


def test_field_not_flagged_as_cumulative_passes_through_untouched(tmp_path):
    sheet = _sheet_plan(write_col=3, populated_period_cols=[2])
    path = _write_workbook(tmp_path, {2: 100.0})
    plan = _plan(path, sheet)

    values, issues = resolve_cumulative_periods(
        plan, str(path), {FIELD: 250.0, "revenue_from_operations": 1000.0}, {FIELD: 6},
        None, _meta(),
    )
    assert values["revenue_from_operations"] == 1000.0
    assert not any(i.field == "revenue_from_operations" for i in issues)
