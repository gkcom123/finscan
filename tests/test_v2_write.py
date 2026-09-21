"""Stage 4 (finscan2.write) — the three invariants, on a real workbook.

Builds a small model in a temp directory rather than mocking openpyxl: the things
that go wrong here — a formula overwritten, a source file modified, a date not
stamped — only go wrong against a real file.
"""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from finscan2.match.schema import ResolvedValue, Source, Values
from finscan2.model.discover import RowLayout, SheetLayout
from finscan2.model.load import resolve
from finscan2.model.schema import ModelMap, RowKey, RowSpec
from finscan2.write.column import AUDIT_SHEET, write_column

BLUE = Font(color="FF0000FF")


def _workbook(path, *, formula_in_target: bool = False):
    """Two dated columns (C, D) plus an empty reserved column E."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    for index, iso in enumerate(("2026-03-31", "2026-06-30")):
        sheet.cell(2, 3 + index, datetime.fromisoformat(iso))
    for row, label, values in ((5, "Revenue", [100, 110]),
                               (6, "Cost of Sales", [-60, -66])):
        sheet.cell(row, 2, label)
        for index, value in enumerate(values):
            cell = sheet.cell(row, 3 + index, value)
            cell.font = BLUE
        sheet.cell(row, 5).font = BLUE
    sheet.cell(7, 2, "Gross Profit")
    sheet.cell(7, 3, "=C5+C6")
    sheet.cell(7, 4, "=D5+D6")
    if formula_in_target:
        sheet.cell(5, 5, "=D5*1.1")
    book.save(path)
    return path


def _map_and_layout(write_col: int = 5):
    rows = [(5, "Revenue", "income_statement", "input", False),
            (6, "Cost of Sales", "income_statement", "input", False),
            (7, "Gross Profit", "income_statement", "formula", True)]
    layout = SheetLayout(
        sheet="Model", label_col=2, header_row=2, first_data_row=5,
        period_dates={3: "2026-03-31", 4: "2026-06-30"}, period_date_row=2,
        reference_col=4, write_col=write_col, units="millions", cadence_months=3,
        rows=[RowLayout(row=r, label=label, section=section, role=role,
                        has_formula=has_formula,
                        formula="=D5+D6" if has_formula else None)
              for r, label, section, role, has_formula in rows],
    )
    model = ModelMap(
        company="acme", sheet="Model", units="millions", cadence_months=3,
        label_col=2, header_row=2, first_data_row=5, reference_col=4,
        write_col=write_col, write_mode="fill_blank",
        period_dates={3: "2026-03-31", 4: "2026-06-30"}, period_date_row=2,
        confirmed_by="test",
        rows=[RowSpec(key=RowKey(label=label, section=section), row_hint=r,
                      kind="formula" if has_formula else "input",
                      resolve=None if has_formula else f"pdf:{label}")
              for r, label, section, role, has_formula in rows],
    )
    return model, layout


def _values(*, blank_row: int | None = None) -> Values:
    values = Values(company="acme", sheet="Model", period_end="2026-09-30",
                    write_col=5, units="millions")
    source = Source(statement="income_statement", page=4, caption="Revenue",
                    column_header="July - September 2026", column_index=0,
                    months=3, end="2026-09-30", printed=120.0)
    values.values.append(ResolvedValue(row=5, label="Revenue",
                                       section="income_statement", value=120.0,
                                       resolve="pdf:Revenue", source=source))
    if blank_row == 6:
        values.values.append(ResolvedValue(
            row=6, label="Cost of Sales", section="income_statement", value=None,
            resolve="pdf:Cost of Sales", unresolved="no row captioned like 'Cost of Sales'"))
    else:
        values.values.append(ResolvedValue(row=6, label="Cost of Sales",
                                           section="income_statement", value=-70.0,
                                           resolve="pdf:Cost of Sales"))
    return values


def _run(tmp_path, *, formula_in_target: bool = False, blank_row: int | None = None):
    source = _workbook(tmp_path / "model.xlsx", formula_in_target=formula_in_target)
    model, layout = _map_and_layout()
    out = tmp_path / "model_updated.xlsx"
    result = write_column(_values(blank_row=blank_row), resolve(model, layout),
                          str(source), str(out))
    return source, out, result


def test_values_land_in_the_write_column(tmp_path):
    _, out, result = _run(tmp_path)
    sheet = load_workbook(out)["Model"]
    assert sheet.cell(5, 5).value == 120.0
    assert sheet.cell(6, 5).value == -70.0
    assert result.values_written == 2


def test_the_source_workbook_is_never_modified(tmp_path):
    source, _, _ = _run(tmp_path)
    sheet = load_workbook(source)["Model"]
    assert sheet.cell(5, 5).value is None
    assert AUDIT_SHEET not in load_workbook(source).sheetnames


def test_a_formula_row_is_copied_forward_with_translated_references(tmp_path):
    _, out, result = _run(tmp_path)
    sheet = load_workbook(out)["Model"]
    assert sheet.cell(7, 5).value == "=E5+E6"
    assert result.formulas_copied == 1


def test_a_formula_already_in_the_target_is_never_overwritten(tmp_path):
    """Checked against the live cell, not the map: a cell an analyst turned into
    a formula after the map was frozen must win."""
    _, out, result = _run(tmp_path, formula_in_target=True)
    sheet = load_workbook(out)["Model"]
    assert sheet.cell(5, 5).value == "=D5*1.1"
    assert any(i.code == "formula_protected" for i in result.issues)


def test_every_written_cell_carries_its_provenance(tmp_path):
    _, out, _ = _run(tmp_path)
    comment = load_workbook(out)["Model"].cell(5, 5).comment
    assert comment is not None
    assert "July - September 2026" in comment.text
    assert "printed 120.00" in comment.text


def test_a_blank_row_is_annotated_with_the_reason(tmp_path):
    """An empty cell with no explanation is indistinguishable from one nobody
    got to."""
    _, out, result = _run(tmp_path, blank_row=6)
    cell = load_workbook(out)["Model"].cell(6, 5)
    assert cell.value is None
    assert "no value written" in cell.comment.text
    assert result.blanks_annotated == 1


def test_the_new_column_is_stamped_with_its_date(tmp_path):
    """Without this, next quarter's run cannot see the column at all — both the
    continuity check and de-cumulation find their columns by date."""
    _, out, _ = _run(tmp_path)
    assert load_workbook(out)["Model"].cell(2, 5).value == datetime(2026, 9, 30)


def test_a_missing_date_row_is_reported_not_ignored(tmp_path):
    source = _workbook(tmp_path / "model.xlsx")
    model, layout = _map_and_layout()
    model.period_date_row = None
    result = write_column(_values(), resolve(model, layout), str(source),
                          str(tmp_path / "out.xlsx"))
    assert any(i.code == "period_not_stamped" for i in result.issues)


def test_the_audit_sheet_records_the_run(tmp_path):
    _, out, _ = _run(tmp_path, blank_row=6)
    book = load_workbook(out)
    assert AUDIT_SHEET in book.sheetnames
    text = "\n".join(str(cell.value) for row in book[AUDIT_SHEET].iter_rows()
                     for cell in row if cell.value is not None)
    assert "acme" in text and "2026-09-30" in text
    assert "no row captioned" in text


def test_no_write_column_refuses_rather_than_guessing(tmp_path):
    source = _workbook(tmp_path / "model.xlsx")
    model, layout = _map_and_layout()
    model.write_col = None
    result = write_column(_values(), resolve(model, layout), str(source),
                          str(tmp_path / "out.xlsx"))
    assert result.errors and result.values_written == 0


def test_a_formula_referencing_no_cell_is_never_copied_forward(tmp_path):
    """Almarai's tax row holds "=-26.573-10.578" — two hand-typed figures. Copying
    it forward puts last quarter's tax in this quarter and calls it a formula copy."""
    from finscan2.model.discover import RowLayout, SheetLayout
    from finscan2.model.schema import ModelMap, RowKey, RowSpec

    source = _workbook(tmp_path / "model.xlsx")
    book = load_workbook(source)
    book["Model"].cell(8, 2, "Zakat and Income Tax")
    book["Model"].cell(8, 4, "=-26.573-10.578")
    book.save(source)

    layout = SheetLayout(
        sheet="Model", label_col=2, header_row=2, first_data_row=5,
        period_dates={3: "2026-03-31", 4: "2026-06-30"}, period_date_row=2,
        reference_col=4, write_col=5, units="millions", cadence_months=3,
        rows=[RowLayout(row=8, label="Zakat and Income Tax", section="income_statement",
                        role="formula", has_formula=True, formula="=-26.573-10.578")])
    model = ModelMap(
        company="acme", sheet="Model", units="millions", cadence_months=3,
        label_col=2, header_row=2, first_data_row=5, reference_col=4, write_col=5,
        period_dates={3: "2026-03-31", 4: "2026-06-30"}, period_date_row=2,
        confirmed_by="test",
        rows=[RowSpec(key=RowKey(label="Zakat and Income Tax", section="income_statement"),
                      row_hint=8, kind="formula")])

    out = tmp_path / "out.xlsx"
    result = write_column(_values(), resolve(model, layout), str(source), str(out))
    cell = load_workbook(out)["Model"].cell(8, 5)
    assert cell.value is None
    assert "typed expression" in cell.comment.text
    assert any(i.code == "literal_formula_not_copied" for i in result.issues)
    assert result.formulas_copied == 0
