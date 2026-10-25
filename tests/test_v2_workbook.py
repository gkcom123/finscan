"""Workbooks Excel opens but openpyxl rejects: unreadable print titles."""
from __future__ import annotations

import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from finscan.model.discover import discover_sheet
from finscan.model.workbook import open_workbook


def _with_bad_print_titles(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 10
    plain = tmp_path / "plain.xlsx"
    book.save(plain)

    broken = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(broken, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = data.replace(
                    b"</workbook>",
                    b'<definedNames><definedName name="_xlnm.Print_Titles" '
                    b'localSheetId="0">#N/A</definedName></definedNames></workbook>')
            target.writestr(item, data)
    return plain, broken


def test_unreadable_print_titles_no_longer_stop_the_workbook_opening(tmp_path):
    _, broken = _with_bad_print_titles(tmp_path)
    with pytest.raises(ValueError):
        load_workbook(broken)

    book = open_workbook(broken)
    try:
        assert book["Model"]["B1"].value == 10
    finally:
        book.close()


def test_a_normal_workbook_is_opened_exactly_as_before(tmp_path):
    plain, _ = _with_bad_print_titles(tmp_path)
    book = open_workbook(plain, data_only=False)
    try:
        assert book.sheetnames == ["Model"]
    finally:
        book.close()


def test_the_lenient_reading_is_not_left_switched_on(tmp_path):
    _, broken = _with_bad_print_titles(tmp_path)
    open_workbook(broken).close()
    with pytest.raises(ValueError):
        load_workbook(broken)


def test_other_failures_are_not_hidden(tmp_path):
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(Exception):
        open_workbook(corrupt)


def test_layout_discovery_reads_a_workbook_with_unreadable_print_titles(tmp_path):
    _, broken = _with_bad_print_titles(tmp_path)
    assert discover_sheet(str(broken), "Model").sheet == "Model"
