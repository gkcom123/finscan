"""A filing published as an Excel workbook (one statement per sheet) read as statements."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from finscan.company import resolve
from finscan.pdf.read import read_pdf
from finscan.pdf.sheets import is_filing_workbook, read_workbook


def _filing(path):
    book = Workbook()

    bs = book.active
    bs.title = "FS conso 2-4"
    bs["A1"], bs["A2"], bs["A3"] = "Acme Public Company Limited", "Statement of Financial Position", "As at 30 June 2026"
    bs["D6"] = "Consolidated financial information"
    bs.merge_cells("D6:G6")
    bs["D7"], bs["F7"] = "Unit: Thousand US Dollar", "Unit: Thousand Baht"
    bs.merge_cells("D7:E7")
    bs.merge_cells("F7:G7")
    for col, text in zip("DEFG", ("30 June", "31 December", "30 June", "31 December")):
        bs[f"{col}8"] = text
    bs["B9"] = "Notes"
    for col, year in zip("DEFG", (2026, 2025, 2026, 2025)):
        bs[f"{col}9"] = year
    bs["A11"] = "Assets"
    bs["A12"], bs["B12"] = "Cash and cash equivalents", 8
    for col, value in zip("DEFG", (100, 90, 3300, 3000)):
        bs[f"{col}12"] = value
    bs["A13"] = "Rights to receive reimbursement from"
    bs["A14"] = "   decommissioning funds"
    for col, value in zip("DEFG", (5, 4, 160, 130)):
        bs[f"{col}14"] = value
    bs["A15"] = "Total assets"
    for col, value in zip("DEFG", (105, 94, 3460, 3130)):
        bs[f"{col}15"] = value

    pl = book.create_sheet("PL-Conso 8-9 (3month)")

    def title(first_row, statement):
        pl[f"A{first_row}"] = "Acme Public Company Limited"
        pl[f"A{first_row + 1}"] = statement
        pl[f"A{first_row + 2}"] = "For the three-month period ended 30 June 2026"
        pl[f"F{first_row + 5}"] = "Consolidated financial information (Unaudited)"
        pl.merge_cells(f"F{first_row + 5}:I{first_row + 5}")
        pl[f"F{first_row + 6}"], pl[f"H{first_row + 6}"] = "Unit: Thousand US Dollar", "Unit: Thousand Baht"
        pl.merge_cells(f"F{first_row + 6}:G{first_row + 6}")
        pl.merge_cells(f"H{first_row + 6}:I{first_row + 6}")
        pl[f"D{first_row + 7}"] = "Note"
        for col, year in zip("FGHI", (2026, 2025, 2026, 2025)):
            pl[f"{col}{first_row + 7}"] = year

    title(1, "Statement of Income")
    pl["A10"] = "Revenues"
    pl["B11"] = "Revenue from sales"
    for col, value in zip("FGHI", (1000, 800, 33000, 26000)):
        pl[f"{col}11"] = value
    pl["B12"] = "Other income"                       # a heading with children
    pl["C13"] = "Interest income"
    for col, value in zip("FGHI", (30, 50, 990, 1600)):
        pl[f"{col}13"] = value
    pl["A19"] = "The condensed notes are an integral part of this information."
    title(21, "Statement of Income")                 # the same statement, next page
    pl["A30"] = "Profit for the period"
    for col, value in zip("FGHI", (400, 300, 13000, 9700)):
        pl[f"{col}30"] = value
    pl["A39"] = "The condensed notes are an integral part of this information."
    title(41, "Statement of Comprehensive Income")   # another statement, same sheet
    pl["A50"] = "Total comprehensive income"
    for col, value in zip("FGHI", (410, 310, 13300, 10000)):
        pl[f"{col}50"] = value
    book.save(path)
    return path


def _model(path):
    book = Workbook()
    book.active.title = "Model"
    book.active["A1"] = "Revenue"
    book.save(path)
    return path


@pytest.fixture
def doc(tmp_path):
    return read_workbook(_filing(tmp_path / "FINANCIAL_STATEMENTS.XLSX"), "sha")


def test_each_statement_is_read_with_its_scope_and_kind(doc):
    assert [(s.kind, s.title) for s in doc.statements] == [
        ("balance_sheet", "Consolidated Statement of Financial Position As at 30 June 2026"),
        ("income_statement",
         "Consolidated Statement of Income For the three-month period ended 30 June 2026"),
        ("comprehensive_income",
         "Consolidated Statement of Comprehensive Income For the three-month period "
         "ended 30 June 2026")]
    assert doc.units == "thousands"


def test_the_notes_column_is_not_a_figure_column_and_merged_headers_reach_each_column(doc):
    balance = doc.statements[0]
    assert len(balance.columns) == 4
    assert [("US Dollar" in c.header, "Baht" in c.header) for c in balance.columns] == [
        (True, False), (True, False), (False, True), (False, True)]
    assert [(c.kind, c.months, c.end) for c in balance.columns[:2]] == [
        ("point_in_time", None, "2026-06-30"), ("point_in_time", None, "2025-12-31")]
    cash = balance.rows[0]
    assert (cash.caption, cash.values) == ("Cash and cash equivalents", [100, 90, 3300, 3000])


def test_a_year_header_takes_its_span_and_end_from_the_title(doc):
    income = doc.statements[1]
    assert [(c.kind, c.months, c.end) for c in income.columns] == [
        ("period", 3, "2026-06-30"), ("period", 3, "2025-06-30"),
        ("period", 3, "2026-06-30"), ("period", 3, "2025-06-30")]


def test_a_statement_printed_over_two_pages_is_one_statement(doc):
    income = doc.statements[1]
    assert [r.caption for r in income.rows] == [
        "Revenue from sales", "Interest income", "Profit for the period"]
    assert income.rows[1].heading == "Other income"


def test_a_wrapped_caption_is_joined(doc):
    captions = [r.caption for r in doc.statements[0].rows]
    assert "Rights to receive reimbursement from decommissioning funds" in captions


def test_a_workbook_of_statements_is_told_from_a_model(tmp_path):
    assert is_filing_workbook(_filing(tmp_path / "f.xlsx"))
    assert not is_filing_workbook(_model(tmp_path / "m.xlsx"))


def test_read_pdf_hands_a_workbook_to_the_sheet_reader(tmp_path):
    doc = read_pdf(_filing(tmp_path / "f.xlsx"))
    assert any(i.code == "workbook_filing" for i in doc.issues)


def test_a_workbook_with_no_statement_is_refused(tmp_path):
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    book.active["A1"] = "nothing here"
    book.save(path)
    assert any(i.code == "no_statements" for i in read_workbook(path, "x").issues)


def test_a_filing_workbook_beside_the_model_is_a_filing_not_a_second_model(tmp_path):
    folder = tmp_path / "inbox" / "Acme"
    folder.mkdir(parents=True)
    (tmp_path / "models").mkdir()
    _model(folder / "Acme_2Q26.xlsx")
    _filing(folder / "FINANCIAL_STATEMENTS.XLSX")
    paths = resolve("acme", tmp_path)
    assert paths.excel.name == "Acme_2Q26.xlsx"
    assert [p.name for p in paths.pdfs] == ["FINANCIAL_STATEMENTS.XLSX"]


def test_a_lone_workbook_is_the_model(tmp_path):
    folder = tmp_path / "inbox" / "Acme"
    folder.mkdir(parents=True)
    (tmp_path / "models").mkdir()
    _model(folder / "Acme_2Q26.xlsx")
    (folder / "release.pdf").write_text("x", encoding="utf-8")
    assert resolve("acme", tmp_path).excel.name == "Acme_2Q26.xlsx"
