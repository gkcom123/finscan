"""The blue/black convention is the whole safety model, so it is tested hard."""
from __future__ import annotations

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.styles.colors import Color

from finscan.excel.style_probe import (
    Role,
    Theme,
    apply_tint,
    classify_color,
    probe_cell,
)


def test_blues_from_different_palettes_all_read_as_input():
    for rgb in ["0000FF", "0070C0", "1F4E79", "305496", "4472C4", "2F5597", "8FAADC"]:
        assert classify_color(rgb)[0] is Role.input, rgb


def test_black_grey_and_absent_colour_read_as_formula():
    for rgb in ["000000", "1A1A1A", "808080", "404040", None]:
        assert classify_color(rgb)[0] is Role.formula, rgb


def test_green_and_red_are_links_not_inputs():
    assert classify_color("00B050")[0] is Role.link_sheet
    assert classify_color("008000")[0] is Role.link_sheet
    assert classify_color("FF0000")[0] is Role.link_external
    assert classify_color("C00000")[0] is Role.link_external


def test_desaturated_blue_grey_heading_is_not_an_input():
    """Office 'Text 2' (#44546A) is a heading colour. Writing to it would
    overwrite section titles, so it must not classify as blue."""
    assert classify_color("44546A")[0] is not Role.input


def test_tint_moves_lightness_only():
    assert apply_tint("4472C4", 0.4) != "4472C4"
    assert apply_tint("000000", 0.35) == "595959"
    assert apply_tint("4472C4", 0.0) == "4472C4"


def test_theme_colour_is_resolved_through_a_saved_workbook(tmp_path):
    """A theme-coloured font has no rgb attribute at all; without theme
    resolution these cells look colourless and would be treated as formulas."""
    p = tmp_path / "themed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 10
    ws["A1"].font = Font(color=Color(theme=8, tint=-0.2))   # accent5 = a blue
    ws["A2"] = 20
    ws["A2"].font = Font(color="000000")
    wb.save(p)

    wb2 = load_workbook(p)
    theme = Theme.from_workbook(wb2)
    ws2 = wb2.active
    assert probe_cell(ws2["A1"], theme).role is Role.input
    assert probe_cell(ws2["A2"], theme).role is Role.formula


def test_a_formula_is_never_writable_however_it_is_painted(tmp_path):
    p = tmp_path / "trap.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = "=A1*2"
    ws["A2"].font = Font(color="0000FF")     # blue-painted formula
    wb.save(p)

    wb2 = load_workbook(p)
    st = probe_cell(wb2.active["A2"], Theme.from_workbook(wb2))
    assert st.has_formula
    assert st.role is Role.formula
    assert st.writable is False


def test_named_input_style_is_honoured(tmp_path):
    p = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 5
    ws["A1"].style = "Input"
    wb.save(p)
    wb2 = load_workbook(p)
    assert probe_cell(wb2.active["A1"], Theme.from_workbook(wb2)).role is Role.input
