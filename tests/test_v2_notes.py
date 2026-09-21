"""Note tables, and the severed-digit repair that makes them safe to read.

Every line here is taken verbatim from Almarai's segment note as the PDF text
layer produced it — including the damage, because the damage is the point.
"""
from __future__ import annotations

from finscan2.pdf import repair
from finscan2.pdf.notes import extract_notes, note_blocks
from finscan2.pdf.statements import parse_row

#: Note 10 as pdfplumber gives it: leading digits severed from their numbers.
_SEGMENT_NOTE = """\
10. SEGMENT REPORTING
The Group's principal business activities involve manufacturing and trading.
Dairy Other
and Juice Bakery Protein* Activities* Total
'000 '000 '000 '000 '000
30 June 2026
Revenue 7,937,368 1,408,790 2,265,235 858,086 1 2,469,479
Depreciation and Amortisation (702,893) (97,651) ( 362,076) (93,687) ( 1,256,307)
Profit / (Loss) for the period 9 15,780 2 31,694 2 41,400 (20,256) 1 ,368,618
Total Assets 2 3,726,722 1,960,260 1 3,105,460 4,015,870 4 2,808,312
31 December 2025
Total Assets 22,267,954 1 ,866,423 12,004,313 3 ,828,208 3 9,966,898
Total Liabilities 13,200,048 381,699 5 ,195,665 662,250 1 9,439,662
"""

_SPANS = {"2026-06-30": 6, "2025-06-30": 6}


# --------------------------------------------------------------------------- #
# The repair
# --------------------------------------------------------------------------- #

def test_severed_digit_groups_are_rejoined():
    line = "Profit / (Loss) for the period 9 15,780 2 31,694 2 41,400 (20,256) 1 ,368,618"
    assert parse_row(repair.rejoin(line)).values == \
        [915780.0, 231694.0, 241400.0, -20256.0, 1368618.0]


def test_a_repair_is_only_trusted_when_the_row_adds_up():
    """The check that makes this safe rather than hopeful: the table prints its own
    total, so a correct repair reproduces it and an incorrect one almost never does."""
    assert repair.totals_agree([915780.0, 231694.0, 241400.0, -20256.0, 1368618.0])
    assert not repair.totals_agree([15780.0, 2.0, 31694.0, 2.0, 41400.0, -20256.0,
                                    1.0, 368618.0])


def test_an_intact_line_is_left_alone():
    line = "Revenue 7,937,368 1,408,790 2,265,235 858,086 12,469,479"
    assert repair.rejoin(line) == line
    assert not repair.looks_severed(line)


def test_a_repair_that_does_not_reconcile_is_refused():
    """A plausible-looking number that no arithmetic supports is exactly what must
    not reach a model."""
    line = "Something 1 23,456 9 87,654 5 55,555"
    used, repaired, why = repair.verified(line, lambda l: parse_row(l).values)
    assert used == line and not repaired
    assert "does not add" in why or "neither" in why


# --------------------------------------------------------------------------- #
# Note structure
# --------------------------------------------------------------------------- #

def test_numbered_notes_are_found_by_their_heading():
    blocks = note_blocks(_SEGMENT_NOTE)
    assert [(n, t) for n, t, _ in blocks] == [(10, "SEGMENT REPORTING")]


def test_a_note_yields_one_statement_per_period_it_reports():
    """The period is a ROW heading in a segment table, not a column heading, so one
    note block covers several periods with identical captions."""
    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    assert [s.heading for s in statements] == ["2026-06-30", "2025-12-31"]
    assert all(s.note == 10 for s in statements)
    assert all(s.kind == "other" for s in statements)


def test_the_period_span_comes_from_the_primary_statements():
    """A note reports "the period then ended"; only the primary statements say how
    long that is. Without the span a flow row cannot be de-cumulated."""
    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    assert statements[0].columns[0].months == 6        # 30 June 2026, interim
    assert statements[1].columns[0].months is None     # 31 December 2025


def test_an_unknown_span_is_reported_not_assumed():
    _, issues = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    assert any(i.code == "note_span_unknown" for i in issues)


def test_the_total_column_is_named_by_position_not_by_counting_words():
    """The header's word count and the table's column count disagree whenever a
    segment name runs to two words ("Dairy and Juice")."""
    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    headers = [c.header for c in statements[0].columns]
    assert headers[-1] == "Total" and len(headers) == 5


def test_figures_survive_the_round_trip():
    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    rows = {r.caption: r.values for r in statements[0].rows}
    assert rows["Depreciation and Amortisation"][-1] == -1256307.0
    assert rows["Total Assets"][-1] == 42808312.0
    assert rows["Revenue"][-1] == 12469479.0


# --------------------------------------------------------------------------- #
# Alignment
# --------------------------------------------------------------------------- #

def test_a_short_row_is_right_aligned_and_verified():
    """"- - - (1,222) (1,222)" parses as two values. Left at indices 0 and 1 the
    TOTAL would sit under the first segment."""
    text = _SEGMENT_NOTE.replace(
        "Total Assets 2 3,726,722 1,960,260 1 3,105,460 4,015,870 4 2,808,312",
        "Share of Results of Associate - - - (1,222) ( 1,222)")
    statements, issues = extract_notes(15, text, _SPANS)
    rows = {r.caption: r.values for r in statements[0].rows}
    share = rows.get("Share of Results of Associate")
    assert share is not None and len(share) == 5
    assert share[-1] == -1222.0 and share[0] is None


def test_a_short_row_that_cannot_be_reconciled_is_dropped_with_a_reason():
    text = _SEGMENT_NOTE.replace(
        "Total Assets 2 3,726,722 1,960,260 1 3,105,460 4,015,870 4 2,808,312",
        "Mystery row 111,111 222,222")
    statements, issues = extract_notes(15, text, _SPANS)
    rows = {r.caption for r in statements[0].rows}
    assert "Mystery row" not in rows
    assert any(i.code == "note_row_unaligned" for i in issues)


# --------------------------------------------------------------------------- #
# Reaching a note from the mapping
# --------------------------------------------------------------------------- #

def test_a_note_is_addressed_by_number_and_period():
    """"Total Assets" appears under two periods with different figures; the period
    end is what separates them."""
    from finscan2.schema import PdfDoc

    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    doc = PdfDoc(path="f.pdf", sha256="x", statements=statements)
    assert doc.note(10, "2026-06-30").rows[-1].values[-1] == 42808312.0
    assert doc.note(10, "2025-12-31").rows[-1].values[-1] == 19439662.0
    assert doc.note(11) is None


def test_the_total_column_wins_when_several_match_the_period():
    """Five segment columns share one period, so the period tests cannot separate
    them. A tie with no total is refused rather than resolved by position."""
    from finscan2.match.select import prefer_total, select_column

    statements, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    column, why = select_column(statements[0], "2026-06-30", 3, point_in_time=False)
    assert column.header == "Total" and "total" in why

    nameless = [c for c in statements[0].columns if c.header != "Total"]
    chosen, note = prefer_total(nameless)
    assert chosen is None and "cannot be decided" in note


def test_a_parser_upgrade_reuses_a_cached_vision_transcription(tmp_path):
    """A vision transcription is not reproducible — the same page reads 137,011 one
    time and 137,001 the next. So a parser upgrade must re-parse without asking the
    model again, or figures a reviewer already checked move for no reason."""
    import json

    from finscan2.pdf import read as read_mod

    cache = tmp_path / "cache"
    cache.mkdir()
    pdf = tmp_path / "f.pdf"
    pdf.write_bytes(b"%PDF-1.4 not a real pdf")
    sha = "deadbeef"

    stale = {
        "parser_version": "OLD", "path": str(pdf), "sha256": sha,
        "pages": [{"page": 1, "source": "vision", "chars": 9, "text": "Revenue 137,011",
                   "tables": 0}],
        "statements": [], "issues": [], "ocr": {"pages": [1], "engine": "vision"},
    }
    (cache / f"{sha}.json").write_text(json.dumps(stale), encoding="utf-8")

    asked = []

    def _never(path, page_no, dpi=300):
        asked.append(page_no)
        return "Revenue 999,999"        # a different reading, as vision does

    # The PDF cannot be opened, so reaching pdfplumber at all would raise; the point
    # is that the cached text is what a re-parse starts from.
    original_sha, original_transcribe = read_mod.sha256_of, read_mod.ocr.transcribe
    read_mod.sha256_of = lambda p: sha
    read_mod.ocr.transcribe = _never
    try:
        try:
            read_mod.read_pdf(pdf, cache_dir=cache)
        except Exception:
            pass          # pdfplumber will refuse the fake file; that is expected
    finally:
        read_mod.sha256_of, read_mod.ocr.transcribe = original_sha, original_transcribe

    assert asked == [], "the model was asked to transcribe a page the cache already had"


def test_learn_finds_a_caption_the_primary_statement_never_prints():
    """Almarai's D&A for the period appears only in the segment note, so a row
    looking for it in the income statement can never match — and stayed blank."""
    from finscan2.model import propose_captions as pc
    from finscan2.model.schema import ModelMap, RowKey, RowSpec
    from finscan2.schema import Column, PdfDoc, Statement, StatementRow

    income = Statement(
        page=6, kind="income_statement", title="t", heading="h",
        columns=[Column(index=0, header="April - June 2026", months=3,
                        end="2026-06-30", kind="period")],
        rows=[StatementRow(caption="Revenue", values=[1.0], raw="")])
    notes, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    doc = PdfDoc(path="f.pdf", sha256="x", statements=[income, *notes])

    model = ModelMap(company="almarai", sheet="Model", rows=[
        RowSpec(key=RowKey(label="Depreciation & amortisation",
                           section="income_statement"),
                row_hint=45, kind="input", statement="income_statement",
                resolve="pdf:Depreciation & amortisation")])
    proposals = pc.deterministic(model, doc)
    assert proposals[0].method == "note"
    assert proposals[0].statement == "note:10"
    assert proposals[0].caption == "Depreciation and Amortisation"

    pc.apply(model, proposals)
    assert model.rows[0].statement == "note:10"
    assert "in the notes" in model.rows[0].review


def test_a_caption_printed_in_two_different_notes_is_refused():
    """The same caption in one note across several periods is not ambiguity — a note
    reports several periods. Two DIFFERENT notes printing it is."""
    from finscan2.model import propose_captions as pc
    from finscan2.model.schema import ModelMap, RowKey, RowSpec
    from finscan2.schema import Column, PdfDoc, Statement, StatementRow

    def _note_statement(number: int) -> Statement:
        return Statement(
            page=15, kind="other", note=number, title=f"Note {number}: X",
            heading="2026-06-30",
            columns=[Column(index=0, header="Total", months=6, end="2026-06-30",
                            kind="period")],
            rows=[StatementRow(caption="Depreciation and Amortisation", values=[1.0],
                               raw="")])

    income = Statement(page=6, kind="income_statement", title="t", heading="h",
                       columns=[], rows=[])
    doc = PdfDoc(path="f.pdf", sha256="x",
                 statements=[income, _note_statement(10), _note_statement(12)])
    model = ModelMap(company="x", sheet="Model", rows=[
        RowSpec(key=RowKey(label="Depreciation and Amortisation",
                           section="income_statement"),
                row_hint=45, kind="input", statement="income_statement",
                resolve="pdf:Depreciation and Amortisation")])
    proposal = pc.deterministic(model, doc)[0]
    assert proposal.caption is None
    assert "note 10, note 12" in proposal.why


def test_a_cumulative_note_figure_becomes_the_quarter(tmp_path):
    """The whole point: the note prints six months, the model wants the quarter, so
    the prior quarter already in the sheet is subtracted."""
    from datetime import datetime

    from openpyxl import Workbook

    from finscan2.mapping.compile import compile_mapping
    from finscan2.mapping.schema import Mapping, MappingRow
    from finscan2.match.resolve import resolve_values
    from finscan2.model.discover import discover_sheet
    from finscan2.model.load import resolve as bind_map
    from finscan2.schema import PdfDoc

    book = Workbook()
    sheet = book.active
    sheet.title = "Model"
    sheet.cell(1, 2, "SAR thousands")
    for index, iso in enumerate(("2025-12-31", "2026-03-30")):
        sheet.cell(2, 3 + index, datetime.fromisoformat(iso))
    sheet.cell(5, 2, "Depreciation & amortisation")
    sheet.cell(5, 3, -600000.0)
    sheet.cell(5, 4, -620000.0)          # Q1 2026, the quarter to subtract
    path = tmp_path / "m.xlsx"
    book.save(path)

    notes, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    doc = PdfDoc(path="f.pdf", sha256="x", units="thousands", statements=notes)
    mapping = Mapping(company="x", sheet="Model", rows=[
        MappingRow(label="Depreciation & amortisation", section="income_statement",
                   statement="note:10", pdf="Depreciation and Amortisation")])

    layout = discover_sheet(str(path), "Model")
    model, issues = compile_mapping(mapping, layout)
    assert not issues
    values = resolve_values(doc, bind_map(model, layout), str(path), "2026-06-30")

    value = values.values[0]
    # H1 -1,256,307 less Q1 -620,000 = Q2 -636,307
    assert value.value == -636307.0
    assert value.source.months == 6
    assert "de-cumulated" in value.adjustments[0].detail


def test_a_primary_statement_is_preferred_over_a_note():
    """Almarai prints D&A in both the cash flow and the segment note, and the two
    differ in SIGN: the cash flow adds it back (+1,256,307), which is the convention
    the model stores, while the note shows it as an expense (-1,256,307). Reaching
    for the note first hands the model the wrong sign."""
    from finscan2.model import propose_captions as pc
    from finscan2.model.schema import ModelMap, RowKey, RowSpec
    from finscan2.schema import Column, PdfDoc, Statement, StatementRow

    income = Statement(page=6, kind="income_statement", title="t", heading="h",
                       columns=[], rows=[StatementRow(caption="Revenue", values=[1.0],
                                                      raw="")])
    cash_flow = Statement(
        page=9, kind="cash_flow", title="t", heading="h",
        columns=[Column(index=0, header="January - June 2026", months=6,
                        end="2026-06-30", kind="period")],
        rows=[StatementRow(caption="Depreciation and Amortisation",
                           values=[1256307.0], raw="")])
    notes, _ = extract_notes(15, _SEGMENT_NOTE, _SPANS)
    doc = PdfDoc(path="f.pdf", sha256="x", statements=[income, cash_flow, *notes])

    model = ModelMap(company="almarai", sheet="Model", rows=[
        RowSpec(key=RowKey(label="Depreciation & amortisation",
                           section="income_statement"),
                row_hint=45, kind="input", statement="income_statement",
                resolve="pdf:Depreciation & amortisation")])
    proposal = pc.deterministic(model, doc)[0]
    assert proposal.method == "restated"
    assert proposal.statement == "cash_flow"

    pc.apply(model, proposal and [proposal])
    assert model.rows[0].statement == "cash_flow"
    assert "sign convention" in model.rows[0].review
