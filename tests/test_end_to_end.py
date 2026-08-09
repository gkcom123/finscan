"""Full graph runs across three unrelated model layouts, LLM stubbed."""
from __future__ import annotations

from openpyxl import load_workbook

from finscan.graph.build import run as run_graph

NORTHWIND = "northwind_industries"
KEY_ROWS = {          # row -> (canonical field, expected value in crores)
    5: ("revenue_from_operations", 1284.50),
    6: ("other_income", 31.20),
    8: ("cost_of_materials", 523.10),
    9: ("purchases_of_stock_in_trade", 62.40),
    10: ("changes_in_inventories", -11.80),
    11: ("employee_benefit_expense", 217.60),
    12: ("finance_costs", 43.10),
    13: ("depreciation_amortisation", 89.20),
    14: ("other_expenses", 184.70),
    17: ("exceptional_items", 0.00),
    19: ("tax_expense", 52.10),
    21: ("eps_basic", 31.31),
    22: ("eps_diluted", 31.14),
}
FORMULA_ROWS = {7: "=F5+F6", 15: "=SUM(F8:F14)", 18: "=F7-F15+F17", 20: "=F18-F19"}


# --------------------------------------------------------------------------- #
# The confirmation gate
# --------------------------------------------------------------------------- #
def test_a_new_company_is_held_for_one_off_confirmation(workspace, stub_llm):
    before = workspace["northwind"].read_bytes()
    state = run_graph(str(workspace["pdf"]), str(workspace["northwind"]),
                      output_path=str(workspace["northwind_out"]), use_llm_mapping=False,
                      require_confirmation=True)

    assert state["status"] == "awaiting_confirmation"
    assert state.get("write_result") is not None
    assert state["write_result"].workbook_path == str(workspace["northwind_out"])
    assert not state["write_result"].sheets
    assert workspace["northwind"].read_bytes() == before
    assert workspace["northwind_out"].exists()
    # ...but a full proposal is still available to review
    assert state["mappings"] and "awaiting_confirmation" in state["report"]


def test_confirmed_company_runs_straight_through(workspace, stub_llm, confirmed):
    run_graph(str(workspace["pdf"]), str(workspace["northwind"]), use_llm_mapping=False,
              dry_run=True)
    confirmed(NORTHWIND, sheets=["P&L Summary"])

    state = run_graph(str(workspace["pdf"]), str(workspace["northwind"]),
                      output_path=str(workspace["northwind_out"]), use_llm_mapping=False)
    assert state["status"] == "ok", state["report"]
    assert state["write_result"].values_written == 13


def test_layout_drift_forces_re_confirmation(workspace, stub_llm, confirmed):
    run_graph(str(workspace["pdf"]), str(workspace["northwind"]), use_llm_mapping=False,
              dry_run=True, require_confirmation=True)
    confirmed(NORTHWIND, sheets=["P&L Summary"])

    wb = load_workbook(workspace["northwind"])
    wb["P&L Summary"].insert_rows(9)          # someone adds a line item
    wb.save(workspace["northwind"])

    state = run_graph(str(workspace["pdf"]), str(workspace["northwind"]),
                      output_path=str(workspace["northwind_out"]), use_llm_mapping=False,
                      require_confirmation=True)
    assert state["status"] == "awaiting_confirmation"
    assert state["profile_state"] == "drifted"
    assert any(i.code == "profile_drift" for i in state["issues"])


# --------------------------------------------------------------------------- #
# Writing: values into blue cells, formulas carried forward
# --------------------------------------------------------------------------- #
def _run_confirmed(workspace, confirmed, model: str, key: str, sheets: list[str]):
    """Upload once (proposes the profile), confirm it, upload again for real."""
    run_graph(str(workspace["pdf"]), str(workspace[model]), company=key,
              use_llm_mapping=False, dry_run=True)
    confirmed(key, sheets=sheets)
    return run_graph(str(workspace["pdf"]), str(workspace[model]), company=key,
                     output_path=str(workspace[f"{model}_out"]), use_llm_mapping=False)


def test_values_land_in_the_right_rows_rescaled(workspace, stub_llm, confirmed):
    state = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    ws = load_workbook(state["write_result"].workbook_path)["P&L Summary"]

    assert ws["F4"].value == "Q1 FY2027 (Consol.)"
    for row, (field, want) in KEY_ROWS.items():
        got = ws.cell(row, 6).value
        assert got is not None, f"row {row} ({field}) was left blank"
        assert abs(got - want) < 0.01, f"row {row} ({field}): got {got}, want {want}"


def test_formula_rows_receive_translated_formulas_not_values(workspace, stub_llm, confirmed):
    """The model's own arithmetic is carried forward, so subtotals keep
    recalculating instead of being frozen to whatever the PDF said."""
    state = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    ws = load_workbook(state["write_result"].workbook_path)["P&L Summary"]

    for row, expected in FORMULA_ROWS.items():
        assert ws.cell(row, 6).value == expected, f"row {row}"
    assert state["write_result"].formulas_copied == 5


def test_copied_formulas_compute_the_figures_the_filing_reported(workspace, stub_llm, confirmed):
    """Cross-check: evaluating the copied formulas by hand against the extracted
    inputs must reproduce the subtotals printed in the PDF."""
    state = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    ws = load_workbook(state["write_result"].workbook_path)["P&L Summary"]
    v = {r: ws.cell(r, 6).value for r in KEY_ROWS}

    total_income = v[5] + v[6]
    total_expenses = sum(v[r] for r in (8, 9, 10, 11, 12, 13, 14))
    pbt = total_income - total_expenses + v[17]
    pat = pbt - v[19]

    assert abs(total_income - 1315.70) < 0.01
    assert abs(total_expenses - 1108.30) < 0.01
    assert abs(pbt - 207.40) < 0.01
    assert abs(pat - 155.30) < 0.01


def test_existing_columns_and_out_of_scope_sheets_are_untouched(workspace, stub_llm, confirmed):
    state = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    wb = load_workbook(state["write_result"].workbook_path)
    ws = wb["P&L Summary"]

    assert ws["E5"].value == 1213.00
    assert ws["E7"].value == "=E5+E6"
    assert ws["A5"].value == "Net Sales"
    assert wb["Cover"]["A1"].value.startswith("Northwind")
    assert wb["Assumptions"]["B1"].value == 0.11


def test_a_formula_cell_is_never_overwritten_even_if_aimed_at_directly(workspace):
    """Invariant 1, tested at its weakest point: a stale profile pointing the
    write straight at a populated column must still not destroy a calculation."""
    from finscan.excel.discovery import discover
    from finscan.excel.writer import write_workbook

    plan = discover(workspace["northwind"])
    sheet = plan.sheet("P&L Summary")
    sheet.write_col = sheet.reference_col          # aim at the live Q4 column

    out = workspace["dir"] / "clobbered.xlsx"
    result, issues = write_workbook(plan, {"revenue_from_operations": 999.0},
                                    header="BAD", output_path=str(out))

    ws = load_workbook(out)["P&L Summary"]
    assert ws["E7"].value == "=E5+E6", "the subtotal formula was destroyed"
    assert ws["E15"].value == "=SUM(E8:E14)"
    assert any(i.code == "formula_protected" for i in issues)
    assert result.sheets[0].rows_skipped >= 5


def test_written_cells_carry_provenance_comments(workspace, stub_llm, confirmed):
    state = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    ws = load_workbook(state["write_result"].workbook_path)["P&L Summary"]
    assert "revenue_from_operations" in ws["F5"].comment.text
    assert "formula copied from E7" in ws["F7"].comment.text


# --------------------------------------------------------------------------- #
# Multi-sheet
# --------------------------------------------------------------------------- #
def test_both_statement_tabs_are_written_each_in_its_own_units(
    workspace, stub_llm, confirmed
):
    state = _run_confirmed(workspace, confirmed, "acme", "acme_manufacturing",
                           ["Consolidated P&L", "Standalone P&L"])
    assert state["status"] == "ok", state["report"]
    assert {s.sheet for s in state["write_result"].sheets} == {"Consolidated P&L",
                                                               "Standalone P&L"}

    wb = load_workbook(state["write_result"].workbook_path)
    assert abs(wb["Consolidated P&L"]["H8"].value - 1284.50) < 0.01     # crores
    assert abs(wb["Standalone P&L"]["H8"].value - 128_450.00) < 0.01    # lakhs
    # EPS is per-share on both tabs, so it is never rescaled — the lakhs tab and
    # the crores tab must show the same figure.
    assert abs(wb["Consolidated P&L"]["H24"].value - 31.31) < 0.01
    assert abs(wb["Standalone P&L"]["H24"].value - 31.31) < 0.01


def test_disabling_a_sheet_in_the_profile_excludes_it(workspace, stub_llm, confirmed):
    state = _run_confirmed(workspace, confirmed, "acme", "acme_manufacturing",
                           ["Consolidated P&L"])
    assert [s.sheet for s in state["write_result"].sheets] == ["Consolidated P&L"]
    wb = load_workbook(state["write_result"].workbook_path)
    assert wb["Standalone P&L"]["H8"].value is None


def test_the_segments_tab_is_never_touched(workspace, stub_llm, confirmed):
    state = _run_confirmed(workspace, confirmed, "acme", "acme_manufacturing",
                           ["Consolidated P&L", "Standalone P&L"])
    seg = load_workbook(state["write_result"].workbook_path)["Segments"]
    assert seg["C1"].value is None and seg["C2"].value is None


def test_reserved_blank_column_is_filled_rather_than_appended_beside(
    workspace, stub_llm, confirmed
):
    state = _run_confirmed(workspace, confirmed, "acme", "acme_manufacturing",
                           ["Consolidated P&L"])
    res = next(s for s in state["write_result"].sheets if s.sheet == "Consolidated P&L")
    assert res.write_mode == "fill_blank" and res.column_letter == "H"


# --------------------------------------------------------------------------- #
# The colourless fallback
# --------------------------------------------------------------------------- #
def test_a_workbook_without_the_convention_is_written_but_flagged(
    workspace, stub_llm, confirmed
):
    state = _run_confirmed(workspace, confirmed, "zenith", "zenith_chemicals", ["Quarterly"])
    assert state["status"] == "needs_review"
    assert any(i.code == "no_colour_convention" for i in state["issues"])

    ws = load_workbook(state["write_result"].workbook_path)["Quarterly"]
    assert abs(ws["F4"].value - 1284.50) < 0.01


# --------------------------------------------------------------------------- #
# Repeat uploads and the retry path
# --------------------------------------------------------------------------- #
def test_second_upload_appends_a_further_column(workspace, stub_llm, confirmed):
    first = _run_confirmed(workspace, confirmed, "northwind", NORTHWIND, ["P&L Summary"])
    second = run_graph(str(workspace["pdf"]), first["write_result"].workbook_path,
                       output_path=str(workspace["dir"] / "twice.xlsx"),
                       period_label="Q2 FY2027", use_llm_mapping=False)

    assert second["status"] == "ok", second["report"]
    res = second["write_result"].sheets[0]
    assert res.column_letter == "G"
    ws = load_workbook(second["write_result"].workbook_path)["P&L Summary"]
    assert ws["G4"].value == "Q2 FY2027"
    assert ws["G7"].value == "=G5+G6"
    assert abs(ws["F5"].value - 1284.50) < 0.01     # the first column survives


def test_bad_extraction_retries_then_reports_needs_review(workspace, monkeypatch, confirmed):
    from conftest import Q1_FY27_LAKHS, sample_extraction
    from finscan.schemas import Extraction

    bad = dict(Q1_FY27_LAKHS)
    bad["profit_after_tax"] = 10_120.00          # last year's PAT: wrong column
    calls = {"n": 0}

    class _Stub:
        def __init__(self, schema):
            self.schema = schema

        def invoke(self, _messages):
            if self.schema is Extraction:
                calls["n"] += 1
                return sample_extraction(bad)
            return self.schema(matches=[])

    monkeypatch.setattr("finscan.extract.extractor.structured", lambda s: _Stub(s))
    monkeypatch.setattr("finscan.llm.factory.structured", lambda s: _Stub(s))

    run_graph(str(workspace["pdf"]), str(workspace["northwind"]), use_llm_mapping=False,
              dry_run=True)
    confirmed(NORTHWIND, sheets=["P&L Summary"])
    state = run_graph(str(workspace["pdf"]), str(workspace["northwind"]),
                      output_path=str(workspace["northwind_out"]), use_llm_mapping=False)

    assert calls["n"] == 3, "one dry run + one retried run = 3 extraction calls"
    assert state["status"] == "needs_review"
    assert any(i.code == "identity_pat" for i in state["issues"])
