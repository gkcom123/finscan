"""Open an analyst's workbook, tolerating junk that Excel itself ignores.

Models inherited through years of copy-paste carry defined names whose value is
`#N/A` or `#REF!`, including a `Print_Titles` one. openpyxl refuses to open such a
file at all ("#N/A is not a valid print titles definition"), which stops a company
before any figure is read. Print titles only affect printing, so an unreadable one
is dropped; nothing else about the workbook is changed.
"""
from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.print_settings import PrintTitles


@contextmanager
def _lenient_print_titles():
    original = PrintTitles.from_string.__func__

    def lenient(cls, value):
        try:
            return original(cls, value)
        except ValueError:
            return cls()

    PrintTitles.from_string = classmethod(lenient)
    try:
        yield
    finally:
        PrintTitles.from_string = classmethod(original)


def open_workbook(path: str | Path, **options: Any) -> Workbook:
    """`openpyxl.load_workbook`, retried once with unreadable print titles dropped.

    Any other failure is raised as it was. The retry never silently changes a
    workbook that opens normally.
    """
    try:
        return load_workbook(path, **options)
    except ValueError:
        with _lenient_print_titles():
            return load_workbook(path, **options)
