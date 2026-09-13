"""Correct a value the filing disclosed only as a year-to-date cumulative figure.

Some fields (commonly Depreciation & Amortisation, drawn from the cash-flow statement or
notes) are sometimes printed only as a cumulative figure even in an otherwise quarterly
filing: Q1 states it "for the three months ended", Q2 "for the six months ended" (Q1+Q2
combined), Q3 "nine months" (Q1+Q2+Q3), and so on. The target Excel model stores each
quarter as its own standalone column, so writing the raw cumulative number into a Q2+
column would silently corrupt that column and everything derived from it (EBITDA, the
arithmetic crosscheck against the filing, any analyst ratio built on it).

extractor._detect_months_covered() tags which fields are cumulative and by how many
months (LineItem.months_covered, threaded through by normalize.to_target_units()). This
module does the actual correction, using the workbook's own already-written prior-quarter
columns for the same row — not a second guess at the number, an honest read of what the
model itself already recorded. When that history isn't there, or doesn't hold together,
this refuses to guess: it drops the field so writer.py's existing missing-field fallback
takes over, and flags the run for review, rather than writing a plausible-looking wrong
number.
"""
from __future__ import annotations

from openpyxl import load_workbook

from finscan.excel.discovery import RowPlan, SheetPlan, WorkbookPlan
from finscan.extract.normalize import UNIT_MULTIPLIER
from finscan.periods import CADENCE_DAYS, infer_cadence_days, months_from_cadence_days, parse_date
from finscan.schemas import Issue, PeriodMeta

#: A single period's own spacing must land within this many months of the model's
#: inferred cadence before its columns are trusted as "one clean period apart". Loose on
#: purpose (whole months, not days) — this is sanity-checking already-written columns,
#: not the tighter day-level filing-vs-column check periods.check_continuity() does.
_MONTH_GAP_TOLERANCE = 1


def _is_number(v: object) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _to_base_units(value: float, units: str) -> float:
    """Sheet units -> base units, matching normalize.convert()'s convention. Falls back
    to treating an unrecognised unit as already-base, same leniency to_target_units()
    already shows for an unrecognised target unit."""
    return value * UNIT_MULTIPLIER.get(units, 1.0)


def _single_period_months(sheet: SheetPlan, meta: PeriodMeta) -> int:
    """This model's own quarter length in months, inferred from its column spacing."""
    dates = [d for d in (parse_date(v) for v in sheet.period_dates.values()) if d]
    cadence_days = infer_cadence_days(dates) or CADENCE_DAYS.get(meta.period_type, 91)
    return months_from_cadence_days(cadence_days)


def _authoritative_sheet(
    plan: WorkbookPlan, field: str, enabled_sheets: set[str] | None
) -> tuple[SheetPlan | None, list[SheetPlan]]:
    """The highest-scoring in-scope, enabled sheet where `field` is a writable row, plus
    every other sheet that also maps it (for the multi-sheet-authority note)."""
    candidates = [
        s for s in plan.in_scope
        if (enabled_sheets is None or s.sheet in enabled_sheets)
        and any(rp.field == field and rp.writable for rp in s.rows)
    ]
    if not candidates:
        return None, []
    candidates.sort(key=lambda s: s.score, reverse=True)
    return candidates[0], candidates[1:]


def _row_for_field(sheet: SheetPlan, field: str) -> RowPlan | None:
    return next((rp for rp in sheet.rows if rp.field == field and rp.writable), None)


def _prior_columns(
    sheet: SheetPlan, n: int, single_period_months: int
) -> tuple[list[int] | None, str | None, str | None]:
    """The N populated period columns immediately before the write column.

    Returns (columns, error_code, reason). error_code is None on success; otherwise one
    of "cumulative_period_insufficient_history" (not enough prior columns exist yet) or
    "cumulative_period_ambiguous_history" (they exist, but their own dates don't fit an
    evenly-spaced sequence of single periods, so summing them would be a guess).
    """
    prior = [c for c in sheet.populated_period_cols if c < sheet.write_col]
    if len(prior) < n:
        return None, "cumulative_period_insufficient_history", (
            f"only {len(prior)} prior period column(s) exist on '{sheet.sheet}' before "
            f"the write column; {n} needed to recover the standalone figure"
        )
    chosen = prior[-n:]

    known = {c: parse_date(sheet.period_dates.get(c)) for c in chosen}
    known = {c: d for c, d in known.items() if d}
    if len(known) == len(chosen) and len(chosen) > 1:
        ordered = sorted(known.items())
        for (c1, d1), (c2, d2) in zip(ordered, ordered[1:]):
            gap_months = round((d2 - d1).days / 30.44)
            if abs(gap_months - single_period_months) > _MONTH_GAP_TOLERANCE:
                return None, "cumulative_period_ambiguous_history", (
                    f"columns {c1} and {c2} on '{sheet.sheet}' are {gap_months} month(s) "
                    f"apart, not the expected {single_period_months} — refusing to guess "
                    f"which columns to sum"
                )
    return chosen, None, None


def _read_row_values(excel_path: str, sheet_name: str, row: int, cols: list[int]) -> list[float | None]:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        return [v if _is_number(v := ws.cell(row, c).value) else None for c in cols]
    finally:
        wb.close()


def resolve_cumulative_periods(
    plan: WorkbookPlan,
    excel_path: str,
    values: dict[str, float],
    months_covered: dict[str, int],
    enabled_sheets: set[str] | None,
    meta: PeriodMeta,
) -> tuple[dict[str, float], list[Issue]]:
    """Return (corrected_values, issues). Only touches fields present in months_covered;
    every other field in `values` passes through unchanged."""
    issues: list[Issue] = []
    out = dict(values)

    for field, months in months_covered.items():
        if field not in out:
            continue

        sheet, other_sheets = _authoritative_sheet(plan, field, enabled_sheets)
        if sheet is None:
            # Flagged as cumulative, but nothing in scope actually maps it to a writable
            # row — nothing to correct against, leave the raw extracted value as-is.
            continue
        if other_sheets:
            issues.append(Issue(
                severity="info", code="cumulative_period_multi_sheet_authority", field=field,
                message=f"{field} is a writable row on more than one sheet "
                        f"({sheet.sheet}, {', '.join(s.sheet for s in other_sheets)}); the "
                        f"correction was computed from '{sheet.sheet}''s history and "
                        f"applied everywhere this field is written."))

        single = _single_period_months(sheet, meta)
        ratio = months / single
        n = round(ratio) - 1
        if n < 0 or abs(ratio - round(ratio)) > 0.15:
            issues.append(Issue(
                severity="warning", code="cumulative_period_cadence_mismatch", field=field,
                message=f"{field}: filing states a {months}-month figure, but '{sheet.sheet}' "
                        f"runs on a {single}-month cadence — {months}/{single} is not a clean "
                        f"multiple. Left as extracted; verify against the source PDF."))
            continue
        if n == 0:
            # First period of the fiscal year: the cumulative figure already IS the
            # standalone figure. Falls out of the formula above with no special case.
            continue

        row = _row_for_field(sheet, field)
        cols, err_code, reason = _prior_columns(sheet, n, single)
        if err_code:
            del out[field]
            issues.append(Issue(
                severity="error", code=err_code, field=field, sheet=sheet.sheet,
                message=f"{field}: filing states {months} months, needs {n} prior "
                        f"quarter(s) already in the model to recover the standalone "
                        f"figure, but {reason}. Wrote nothing for this field; needs "
                        f"manual entry."))
            continue

        prior_values = _read_row_values(excel_path, sheet.sheet, row.row, cols)
        if any(v is None for v in prior_values):
            del out[field]
            issues.append(Issue(
                severity="error", code="cumulative_period_insufficient_history",
                field=field, sheet=sheet.sheet,
                message=f"{field}: column(s) {cols} on '{sheet.sheet}' row {row.row} are "
                        f"not all numeric; cannot recover the standalone figure. Wrote "
                        f"nothing for this field; needs manual entry."))
            continue

        prior_sum_base = _to_base_units(sum(prior_values), sheet.units)
        raw = out[field]
        corrected = raw - prior_sum_base
        cell_refs = ", ".join(f"{sheet.sheet}!row{row.row}/col{c}" for c in cols)

        if abs(corrected) > abs(raw):
            issues.append(Issue(
                severity="warning", code="implausible_cumulative_correction", field=field,
                message=f"{field}: corrected standalone-quarter value ({corrected:,.2f}) is "
                        f"larger in magnitude than the raw {months}-month cumulative figure "
                        f"({raw:,.2f}) it was derived from. Written anyway; verify."))

        out[field] = corrected
        issues.append(Issue(
            severity="info", code="cumulative_period_corrected", field=field,
            message=f"{field}: filing stated a {months}-month cumulative figure "
                    f"({raw:,.2f}); subtracted the prior {n} quarter(s) already in "
                    f"'{sheet.sheet}' ({cell_refs}, summing {prior_sum_base:,.2f}) to "
                    f"recover this quarter's standalone value: {corrected:,.2f}."))

    return out, issues
