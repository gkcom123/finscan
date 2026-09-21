"""Read one sheet of an Excel model and describe its layout.

Deterministic and openpyxl-only. This is what `learn` turns into a map; the
quarterly run never calls it, because by then the map already says everything
this would have to work out again.

Scope is one sheet — the Model sheet — by design. Multi-sheet authority is a
later phase, and leaving it out keeps the rules here simple enough to verify.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from finscan2.model.schema import Section

#: Section keywords, checked against a normalized caption. A row inherits the
#: section of the last keyword above it, which is how a cash-flow "Depreciation"
#: is kept apart from the P&L one.
SECTION_KEYWORDS: list[tuple[str, Section]] = [
    ("cash flow", "cash_flow"), ("cashflow", "cash_flow"),
    ("cash flows", "cash_flow"), ("cash flow statement", "cash_flow"),
    ("balance sheet", "balance_sheet"), ("financial position", "balance_sheet"),
    ("income statement", "income_statement"), ("profit and loss", "income_statement"),
    ("profit or loss", "income_statement"), ("p l", "income_statement"),
    ("ltm", "other"), ("credit ratios", "other"), ("credit metrics", "other"),
    ("covenant", "other"), ("valuation", "other"), ("assumptions", "other"),
    ("segment", "other"), ("kpi", "other"), ("guidance", "other"),
    ("capitalisation", "other"), ("capitalization", "other"),
]

#: Rows above any section marker belong to the statement a model opens with.
DEFAULT_SECTION: Section = "income_statement"

#: Generic words that are also valid line-item prefixes; only an exact caption
#: match may flip the section on these.
_EXACT_ONLY = {"valuation", "assumptions", "segment", "kpi", "guidance",
               "covenant", "capitalisation", "capitalization"}

MAX_SCAN_ROWS = 400
MAX_SCAN_COLS = 80

_UNIT_PATTERNS: list[tuple[str, str]] = [
    (r"\bcrores?\b|\bcr\.?\b", "crores"),
    (r"\blakhs?\b|\blacs?\b", "lakhs"),
    (r"\bbillions?\b|\bbn\b", "billions"),
    (r"\bmillions?\b|\bmn\b|\bmm\b", "millions"),
    (r"['’]000s?\b|\bthousands?\b", "thousands"),
]


@dataclass
class RowLayout:
    row: int
    label: str
    section: Section
    #: The role of this row's cell in the reference column: input, formula, link…
    role: str
    has_formula: bool
    #: True when the reference cell begins with "=" but references no other cell —
    #: "=-26.573-10.578", an analyst adding two figures in place. Excel calls that
    #: a formula; this pipeline must not, because copying it forward would carry
    #: last quarter's numbers into this quarter while reporting a formula copy.
    literal_expression: str | None = None
    formula: str | None = None
    reference_value: float | None = None


@dataclass
class SheetLayout:
    sheet: str
    label_col: int = 2
    header_row: int = 1
    first_data_row: int = 1
    period_cols: list[int] = field(default_factory=list)
    period_dates: dict[int, str] = field(default_factory=dict)
    #: The row the dated period headers were found in. The writer must put
    #: the new column's date there, or next quarter's run cannot see this
    #: column at all — continuity and de-cumulation both read these dates.
    period_date_row: int | None = None
    period_headers: dict[int, str] = field(default_factory=dict)
    reference_col: int | None = None
    write_col: int | None = None
    write_mode: str = "append"
    units: str = "units"
    cadence_months: int = 3
    colors_found: bool = False
    rows: list[RowLayout] = field(default_factory=list)


def normalize(text: str) -> str:
    from finscan2.model.schema import normalize_label

    return normalize_label(text)


def detect_section(label: str) -> Section | None:
    norm = normalize(label)
    if not norm or len(norm) > 60:
        return None
    for keyword, section in SECTION_KEYWORDS:
        if keyword in _EXACT_ONLY:
            if norm == keyword:
                return section
            continue
        if norm == keyword or norm.startswith(keyword + " ") or norm.endswith(" " + keyword):
            return section
    return None


def assign_sections(labels: dict[int, str]) -> dict[int, Section]:
    current: Section = DEFAULT_SECTION
    out: dict[int, Section] = {}
    for row in sorted(labels):
        found = detect_section(labels[row])
        if found:
            current = found
        out[row] = current
    return out


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


#: An A1-style reference, excluding a bare number or a function name.
_CELL_REFERENCE = __import__("re").compile(
    r"(?<![A-Za-z0-9_$.])(?:'[^']+'|\[[^\]]+\])?!?\$?[A-Za-z]{1,3}\$?\d+(?![\w(])")


def _references_cells(formula: str) -> bool:
    """True when a formula reads any other cell, on this sheet or another.

    The test for "is this a calculation" has to be this, not whether the text
    begins with "=": a cell holding "=1555.181+12021.836" computes nothing that
    the model maintains, and treating it as a formula carries stale figures
    forward under a provenance that says "formula copied".
    """
    if not formula.startswith("="):
        return False
    # Strings can contain anything that looks like a reference.
    body = __import__("re").sub(r'"[^"]*"', "", formula)
    return bool(_CELL_REFERENCE.search(body))


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _as_iso(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        from calendar import monthrange
        for fmt in ("%b %Y", "%B %Y"):
            try:
                d = datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
            return date(d.year, d.month, monthrange(d.year, d.month)[1]).isoformat()
    return None


def _caption(worksheet, values_sheet, row: int, col: int) -> str:
    """The caption a reader sees in a cell.

    A model often mirrors a block of line items further down the sheet by making
    the caption itself a formula ("=+B121"). Read undisplayed that is a one-letter
    string, and the row drops out of the map entirely — on Almarai's sheet eight
    leverage and coverage rows disappeared exactly that way. The cached value is
    what the analyst actually sees, so it is used whenever a caption is a formula.
    """
    text = _text(worksheet.cell(row, col).value)
    if text.startswith("="):
        return _text(values_sheet.cell(row, col).value)
    return text


def _pick_label_col(worksheet, values_sheet, max_row: int,
                    max_col: int) -> tuple[int, dict[int, str]]:
    """The column holding the line-item captions.

    Chosen by how many cells read as a caption — at least three letters and not a
    number — rather than by raw text count, which picks a notes column whenever a
    model keeps one beside the labels.
    """
    best: tuple[int, dict[int, str], int] = (1, {}, -1)
    for col in range(1, min(max_col, 8) + 1):
        labels = {}
        for row in range(1, max_row + 1):
            text = _caption(worksheet, values_sheet, row, col)
            if len(re.findall(r"[A-Za-zÀ-ɏ]", text)) >= 3:
                labels[row] = text
        if len(labels) > best[2]:
            best = (col, labels, len(labels))
    return best[0], best[1]


def _sniff_units(worksheet, max_row: int, max_col: int, first_data_row: int) -> str:
    for row in range(1, min(first_data_row + 1, max_row) + 1):
        for col in range(1, max_col + 1):
            text = _text(worksheet.cell(row, col).value).lower()
            for pattern, unit in _UNIT_PATTERNS:
                if re.search(pattern, text):
                    return unit
    return "units"


def _cadence_months(dates: list[str]) -> int:
    """The model's own period length, from the spacing of its dated columns."""
    parsed = sorted(date.fromisoformat(d) for d in dates)
    if len(parsed) < 2:
        return 3
    gaps = [round((b - a).days / 30.44) for a, b in zip(parsed, parsed[1:])]
    gaps = [g for g in gaps if g > 0]
    if not gaps:
        return 3
    gaps.sort()
    median = gaps[len(gaps) // 2]
    return min((3, 6, 12), key=lambda candidate: abs(candidate - median))


def _most_recent(cols: list[int], dates: dict[int, str]) -> int | None:
    """The column holding the latest period, by date rather than by position.

    Position is the wrong test whenever a model keeps an older block to the right
    of the live one — Almarai's sheet carries 2019-dated columns after the current
    quarter, and taking the rightmost populated column puts both the reference and
    the write column in that stale block.
    """
    if not cols:
        return None
    dated = [(dates[c], c) for c in cols if c in dates]
    if dated:
        return max(dated)[1]
    return max(cols)


def discover_sheet(path: str, sheet: str) -> SheetLayout:
    """Describe one sheet's layout. Opens the workbook read-only; writes nothing."""
    from finscan.excel.style_probe import Role, Theme, probe_cell

    workbook = load_workbook(path, data_only=False)
    values_book = load_workbook(path, data_only=True)
    try:
        theme = Theme.from_workbook(workbook)
        worksheet = workbook[sheet]
        values_sheet = values_book[sheet]

        max_row = min(worksheet.max_row or 1, MAX_SCAN_ROWS)
        max_col = min(worksheet.max_column or 1, MAX_SCAN_COLS)

        label_col, labels = _pick_label_col(worksheet, values_sheet, max_row, max_col)
        sections = assign_sections(labels)
        layout = SheetLayout(sheet=sheet, label_col=label_col)

        # First row carrying data. Two cells minimum: a header block often holds a
        # single stray formula (an LTM header such as "=AG2") which would
        # otherwise be read as the start of the body.
        def data_cells(row: int) -> int:
            count = 0
            for col in range(label_col + 1, max_col + 1):
                value = worksheet.cell(row, col).value
                if _is_number(value) or (isinstance(value, str) and value.startswith("=")):
                    count += 1
            return count

        counts = {row: data_cells(row) for row in range(1, max_row + 1)}
        data_rows = [r for r, n in counts.items() if n >= 2] or [r for r, n in counts.items() if n]
        if not data_rows:
            return layout
        layout.first_data_row = min(data_rows)

        header_row, best_fill = max(1, layout.first_data_row - 1), -1
        for row in range(1, layout.first_data_row):
            fill = sum(1 for col in range(label_col + 1, max_col + 1)
                       if _text(worksheet.cell(row, col).value))
            if fill >= best_fill and fill > 0:
                header_row, best_fill = row, fill
        layout.header_row = header_row

        # Column census over the rows that carry line items.
        line_rows = [r for r in labels if r >= layout.first_data_row]
        stats: dict[int, dict[str, int]] = {}
        for col in range(label_col + 1, max_col + 1):
            counters = {"input": 0, "formula": 0, "values": 0}
            for row in line_rows:
                style = probe_cell(worksheet.cell(row, col), theme)
                if style.has_formula:
                    counters["formula"] += 1
                elif style.role is Role.input:
                    counters["input"] += 1
                if worksheet.cell(row, col).value is not None:
                    counters["values"] += 1
            if (counters["values"] or counters["input"] or counters["formula"]
                    or _text(worksheet.cell(header_row, col).value)):
                stats[col] = counters

        if not stats:
            return layout

        layout.period_cols = sorted(stats)
        layout.colors_found = any(s["input"] for s in stats.values())
        layout.period_headers = {
            col: _text(worksheet.cell(header_row, col).value)
            for col in layout.period_cols
            if _text(worksheet.cell(header_row, col).value)
        }

        # Dates live in whichever header row carries real datetimes, which is
        # rarely the row picked as the header (that one often reads "1Q", "FY").
        for row in range(1, layout.first_data_row):
            found = {col: _as_iso(worksheet.cell(row, col).value) for col in layout.period_cols}
            found = {c: d for c, d in found.items() if d}
            if len(found) > len(layout.period_dates):
                layout.period_dates = found
                layout.period_date_row = row

        if layout.period_dates:
            layout.cadence_months = _cadence_months(list(layout.period_dates.values()))

        populated = [c for c in layout.period_cols if stats[c]["values"] > 0]
        last_populated = _most_recent(populated, layout.period_dates)

        # A pre-formatted but empty column is the slot the template already
        # reserves for this quarter; filling it beats appending beside it.
        reserved = [c for c in layout.period_cols
                    if (last_populated is None or c > last_populated)
                    and stats[c]["values"] == 0
                    and (stats[c]["input"] or stats[c]["formula"])]
        if reserved:
            layout.write_col, layout.write_mode = reserved[0], "fill_blank"
        elif last_populated is not None:
            layout.write_col, layout.write_mode = last_populated + 1, "append"
        else:
            layout.write_col, layout.write_mode = layout.period_cols[-1] + 1, "append"

        # Reference column: the most recent populated column with a dense input
        # footprint. Picking merely the last column with any blue cell can land on
        # a formula-heavy analysis column and make every row look unwritable.
        with_inputs = [c for c in populated if stats[c]["input"]]
        if with_inputs:
            densest = max(stats[c]["input"] for c in with_inputs)
            floor = max(3, int(densest * 0.5))
            dense = [c for c in with_inputs if stats[c]["input"] >= floor] or with_inputs
            layout.reference_col = _most_recent(dense, layout.period_dates)
        elif populated:
            layout.reference_col = last_populated

        layout.units = _sniff_units(worksheet, max_row, max_col, layout.first_data_row)

        reference = layout.reference_col or (layout.period_cols[-1] if layout.period_cols else None)
        for row in sorted(labels):
            if row < layout.first_data_row or reference is None:
                continue
            cell = worksheet.cell(row, reference)
            style = probe_cell(cell, theme)
            role = style.role.value

            # A formula with no cell references is an input in disguise: the figures
            # are typed into it, so it must be written afresh each period, never
            # copied. Almarai's "Zakat and Income Tax" is "=-26.573-10.578" — two
            # hand-typed components — and nine rows of that model are like it.
            literal = None
            has_formula = style.has_formula
            if has_formula and not _references_cells(str(cell.value)):
                literal = str(cell.value)
                has_formula = False
                role = Role.input.value
            if not has_formula and not layout.colors_found:
                # No colour convention anywhere: a non-formula cell is the best
                # available guess at an input, and the sheet is flagged so a
                # human reviews every row rather than trusting the guess.
                role = Role.input.value
            cached = values_sheet.cell(row, reference).value
            layout.rows.append(RowLayout(
                row=row, label=labels[row], section=sections.get(row, DEFAULT_SECTION),
                role=role, has_formula=has_formula,
                formula=str(cell.value) if has_formula else None,
                literal_expression=literal,
                reference_value=cached if _is_number(cached) else None,
            ))
        return layout
    finally:
        workbook.close()
        values_book.close()


def fingerprint(layout: SheetLayout) -> str:
    """A hash of the sheet's shape, to notice that the model changed.

    Built from captions and their sections, not row numbers: inserting a row
    should not invalidate a map, but renaming or removing line items should.
    """
    import hashlib

    material = "|".join(f"{normalize(r.label)}~{r.section}" for r in layout.rows)
    return hashlib.sha256(material.encode("utf-8")).hexdigest()[:16]
